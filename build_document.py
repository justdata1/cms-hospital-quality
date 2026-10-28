from pathlib import Path
import re
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx import Document
from docx.shared import Inches, Pt, RGBColor
root=Path(__file__).parent
doc=Document()
s=doc.sections[0]
s.top_margin=s.bottom_margin=Inches(.75)
s.left_margin=s.right_margin=Inches(.85)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(5)
for name,size in [('Title',22),('Heading 1',15)]:
 st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(5)
for line in (root/'docs/requirements.md').read_text().splitlines():
 if not line.strip():continue
 if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
 elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
 elif re.fullmatch(r'\[([^]]+)\]\((https://[^)]+)\)',line):
  label,url=re.fullmatch(r'\[([^]]+)\]\((https://[^)]+)\)',line).groups()
  p=doc.add_paragraph()
  h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
  run=OxmlElement('w:r');props=OxmlElement('w:rPr')
  color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');props.append(color)
  underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');props.append(underline)
  run.append(props);t=OxmlElement('w:t');t.text=label;run.append(t);h.append(run);p._p.append(h)
 elif line.startswith('- '):doc.add_paragraph(line[2:],'List Bullet')
 else:doc.add_paragraph(line)
for tree in [doc.styles.element,doc.element]:
 for el in list(tree.iter()):
  if el.tag.endswith('}pBdr'):el.getparent().remove(el)
doc.core_properties.author=''
doc.core_properties.title='CMS Hospital Quality Database Requirements'
doc.save(root/'docs/CMS_Hospital_Quality_Requirements.docx')
